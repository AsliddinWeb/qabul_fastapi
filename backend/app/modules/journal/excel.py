"""Excel builders for the Tizim daftari export.

Two variants (chosen from the dropdown on the page):

  dates       — "Sanalar kesimida": events grouped under a per-day header
                row, plus a "Kunlik xulosa" sheet (day × event-type counts).
  applicants  — "Abituriyentlar kesimida": one block per applicant (header
                row + their full path in order), plus a "Yo'l xulosasi"
                sheet with one row per applicant and the milestone dates.
"""

from __future__ import annotations

import io
from collections import OrderedDict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.modules.journal.service import EVENT_CODES, EVENT_LABELS, STATUS_LABELS

_WEEKDAYS = ["Dushanba", "Seshanba", "Chorshanba", "Payshanba", "Juma", "Shanba", "Yakshanba"]

_HEAD_FILL = PatternFill("solid", fgColor="4F46E5")
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
_GROUP_FILL = PatternFill("solid", fgColor="E0E7FF")
_GROUP_FONT = Font(bold=True, color="1E1B4B", size=11)
_TOTAL_FONT = Font(bold=True)
_THIN = Side(style="thin", color="CBD5E1")
_GROUP_BORDER = Border(top=_THIN, bottom=_THIN)

# Light tint per event type so the eye can scan a long sheet.
_EVENT_FILL = {
    "registered": "F1F5F9",
    "registered_by_staff": "F1F5F9",
    "application_submitted": "FEF3C7",
    "status_changed": "E0F2FE",
    "contract_created": "DCFCE7",
    "contract_renewed": "D1FAE5",
    "contract_signed": "BBF7D0",
    "contract_cancelled": "FFE4E6",
}


def _fmt_dt(d: datetime) -> str:
    return d.strftime("%d.%m.%Y %H:%M")


def _fmt_d(d: datetime | date | None) -> str:
    return d.strftime("%d.%m.%Y") if d else ""


def _header(ws, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.fill = _HEAD_FILL
        c.font = _HEAD_FONT
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.freeze_panes = f"A{ws.max_row + 1}"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _group_row(ws, text: str, ncols: int) -> None:
    ws.append([text] + [""] * (ncols - 1))
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    cell = ws.cell(row=r, column=1)
    cell.fill = _GROUP_FILL
    cell.font = _GROUP_FONT
    cell.alignment = Alignment(vertical="center")
    cell.border = _GROUP_BORDER
    ws.row_dimensions[r].height = 20


def _tint(ws, code: str, ncols: int, event_col: int) -> None:
    color = _EVENT_FILL.get(code)
    if color:
        ws.cell(row=ws.max_row, column=event_col).fill = PatternFill("solid", fgColor=color)
    for c in range(1, ncols + 1):
        ws.cell(row=ws.max_row, column=c).alignment = Alignment(vertical="top", wrap_text=True)


def _title_rows(ws, title: str, period: str, ncols: int) -> None:
    ws.append([title])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.append([period])
    ws.cell(row=2, column=1).font = Font(italic=True, color="64748B")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.append([])


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def period_text(date_from: date | None, date_to: date | None) -> str:
    if not date_from and not date_to:
        return "Davr: barcha vaqt · O'qishni ko'chirish (perevod) kiritilmagan"
    return (f"Davr: {_fmt_d(date_from) or '…'} — {_fmt_d(date_to) or '…'}"
            " · O'qishni ko'chirish (perevod) kiritilmagan")


# ---------------------------------------------------------------- by dates
def build_by_dates(events: list[dict], period: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sanalar kesimida"
    headers = ["№", "Vaqt", "F.I.O", "JSHSHIR", "Telefon", "Hodisa", "Tafsilot",
               "Yo'nalish", "Kim tomonidan"]
    ncols = len(headers)
    _title_rows(ws, "Tizim daftari — sanalar kesimida", period, ncols)
    _header(ws, headers, [6, 8, 34, 17, 15, 24, 48, 36, 28])

    by_day: OrderedDict[date, list[dict]] = OrderedDict()
    for e in events:
        by_day.setdefault(e["at"].date(), []).append(e)

    n = 0
    for day, items in by_day.items():
        _group_row(ws, f"{day.strftime('%d.%m.%Y')} · {_WEEKDAYS[day.weekday()]} — {len(items)} ta hodisa", ncols)
        for e in items:
            n += 1
            ws.append([n, e["at"].strftime("%H:%M"), e["full_name"], e["pinfl"], e["phone"],
                       e["event"], e["detail"], e["program"], e["actor"]])
            _tint(ws, e["code"], ncols, event_col=6)

    # ---- daily summary
    ss = wb.create_sheet("Kunlik xulosa")
    sum_headers = ["Sana", "Hafta kuni"] + [EVENT_LABELS[c] for c in EVENT_CODES] + ["Jami"]
    _title_rows(ss, "Kunlik xulosa", period, len(sum_headers))
    _header(ss, sum_headers, [12, 12] + [16] * len(EVENT_CODES) + [10])
    totals = {c: 0 for c in EVENT_CODES}
    for day, items in by_day.items():
        counts = {c: 0 for c in EVENT_CODES}
        for e in items:
            counts[e["code"]] = counts.get(e["code"], 0) + 1
        for c in EVENT_CODES:
            totals[c] += counts[c]
        ss.append([day.strftime("%d.%m.%Y"), _WEEKDAYS[day.weekday()]]
                  + [counts[c] or "" for c in EVENT_CODES] + [len(items)])
    ss.append(["JAMI", ""] + [totals[c] for c in EVENT_CODES] + [sum(totals.values())])
    for c in ss[ss.max_row]:
        c.font = _TOTAL_FONT
        c.fill = _GROUP_FILL
    return _save(wb)


# ----------------------------------------------------------- by applicants
def build_by_applicants(events: list[dict], period: str) -> bytes:
    wb = Workbook()

    # Group, preserving the service's (name, applicant, time) ordering.
    people: OrderedDict[str, list[dict]] = OrderedDict()
    for e in events:
        people.setdefault(str(e["applicant_id"]), []).append(e)

    # ---- sheet 1: path summary (one row per applicant)
    ws = wb.active
    ws.title = "Yo'l xulosasi"
    headers = ["№", "F.I.O", "JSHSHIR", "Telefon", "Ro'yxatdan o'tgan", "Kim ro'yxatga olgan",
               "Ariza topshirgan", "Yo'nalish", "Ariza holati", "Shartnoma olgan",
               "Shartnoma raqami", "Imzolangan", "Oxirgi yangilanish", "Bekor qilingan",
               "Hodisalar soni"]
    ncols = len(headers)
    _title_rows(ws, "Tizim daftari — abituriyentlar kesimida (xulosa)", period, ncols)
    _header(ws, headers, [6, 34, 17, 15, 14, 26, 14, 36, 18, 14, 18, 14, 16, 14, 10])

    for i, (_, items) in enumerate(people.items(), start=1):
        first = items[0]
        reg = next((e for e in items if e["code"] in ("registered", "registered_by_staff")), None)
        subm = next((e for e in items if e["code"] == "application_submitted"), None)
        status = None
        for e in items:
            if e["code"] == "application_submitted" and status is None:
                status = "topshirildi"
            if e["code"] == "status_changed" and e["to_status"]:
                status = e["to_status"]
        got = next((e for e in items if e["code"] == "contract_created"), None)
        renewed = [e for e in items if e["code"] == "contract_renewed"]
        signed = [e for e in items if e["code"] == "contract_signed"]
        cancelled = [e for e in items if e["code"] == "contract_cancelled"]
        contract_events = [e for e in items if e["contract_number"]]
        last_contract = contract_events[-1]["contract_number"] if contract_events else ""
        ws.append([
            i, first["full_name"], first["pinfl"], first["phone"],
            _fmt_d(reg["at"]) if reg else "",
            ("O'zi" if reg["code"] == "registered" else reg["actor"]) if reg else "",
            _fmt_d(subm["at"]) if subm else "",
            subm["program"] if subm else "",
            STATUS_LABELS.get(status, status or ""),
            _fmt_d(got["at"]) if got else "",
            last_contract,
            _fmt_d(signed[-1]["at"]) if signed else "",
            _fmt_d(renewed[-1]["at"]) if renewed else "",
            _fmt_d(cancelled[-1]["at"]) if cancelled else "",
            len(items),
        ])
        for c in range(1, ncols + 1):
            ws.cell(row=ws.max_row, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    # ---- sheet 2: detailed path per applicant
    ds = wb.create_sheet("Batafsil yo'l")
    d_headers = ["№", "Sana va vaqt", "Hodisa", "Tafsilot", "Yo'nalish", "Kim tomonidan"]
    d_ncols = len(d_headers)
    _title_rows(ds, "Tizim daftari — abituriyentlar kesimida (batafsil)", period, d_ncols)
    _header(ds, d_headers, [6, 17, 24, 52, 36, 30])
    for i, (_, items) in enumerate(people.items(), start=1):
        p = items[0]
        _group_row(ds, f"{i}. {p['full_name']} · JSHSHIR {p['pinfl'] or '—'} · {p['phone'] or '—'} — {len(items)} ta hodisa", d_ncols)
        for j, e in enumerate(items, start=1):
            ds.append([j, _fmt_dt(e["at"]), e["event"], e["detail"], e["program"], e["actor"]])
            _tint(ds, e["code"], d_ncols, event_col=3)
    return _save(wb)
